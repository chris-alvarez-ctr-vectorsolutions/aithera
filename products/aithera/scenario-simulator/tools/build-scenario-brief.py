#!/usr/bin/env python3
"""
Build the SCENARIO BRIEF PowerPoint template — the K&A team's input to the
Scenario Editor — plus a worked example filled with the Marshall scenario.

    python3 tools/build-scenario-brief.py            # writes docs/templates/*.pptx

WHY THIS IS BUILT AND NOT HAND-DRAWN
Every field on every slide is a real PowerPoint PLACEHOLDER on a custom slide
layout (a `<p:ph type="body">` with a fixed idx and a machine-readable shape
name). That is the one property that makes the deck a lossless input:
PowerPoint's Outline view exports title + body placeholder text and nothing
else, so a designer can copy the outline straight into the editor's wizard,
and a future .pptx importer can read fields by placeholder name/idx instead
of guessing from position. Text boxes and tables would silently fall out of
that export — so the layouts carry the decoration (panels, helper text) and
the slides carry only placeholders.

Two decks come out of one field list:
  scenario-brief-template.pptx         labels only, values empty
  scenario-brief-example-marshall.pptx the POC deck v3 content re-homed

Requires python-pptx (1.0.x) and lxml. No other tooling.
"""
import math
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, '..', 'docs', 'templates'))

# ---- palette / type ------------------------------------------------------
INK = '1F2A37'
INK2 = '43505E'
MUTED = '6B7785'
BLUE = '0271CE'      # Vector blue — matches the editor's components
TINT = 'EEF4FB'      # field panel
TINT2 = 'E3ECF7'     # chip
LINE = 'C9D4E0'
WHITE = 'FFFFFF'
TIER = {'unthoughtful': 'B54B3A', 'neutral': '8A6D1F', 'strong': '2E7D4F'}
FONT = 'Calibri'
SLIDE_W, SLIDE_H = 13.333, 7.5
HELP_H = 0.30        # helper strip at the bottom of every panel

TIERS = ['unthoughtful', 'neutral', 'strong']
TIER_TITLE = {'unthoughtful': 'UNTHOUGHTFUL', 'neutral': 'NEUTRAL', 'strong': 'STRONG'}


def rgb(hex6):
    return RGBColor.from_string(hex6)


# =============================================================================
# THE FIELD LIST — this is the template. The .pptx is a rendering of it.
# key      machine name written on the shape (future importer reads this)
# label    the pre-filled first line of the placeholder (Outline export keeps it)
# help     layout-level guidance under the field (never exported)
# box      x, y, w, h in inches (the panel; the placeholder sits above the help)
# size     default body size in pt
# subs     sub-labels written as their own lines inside the placeholder
# =============================================================================
def F(key, label, box, help, size=12, subs=None, tier=None, verbatim=False):
    return dict(key=key, label=label, box=box, help=help, size=size, subs=subs, tier=tier, verbatim=verbatim)


def tier_boxes(y, h, subs, key_prefix, help_text, size=10):
    """Three equal columns across the content width."""
    gap = 0.24
    w = (12.333 - 2 * gap) / 3
    out = []
    for i, t in enumerate(TIERS):
        x = 0.5 + i * (w + gap)
        out.append(F(f'{key_prefix}.{t}', TIER_TITLE[t], (x, y, w, h), help_text, size=size, subs=subs, tier=t))
    return out


LAYOUTS = {
    'Cover': dict(
        title='SCENARIO BRIEF', chip='Cover · the team’s own record — not imported',
        fields=[
            F('cover.title', 'Scenario title', (0.5, 1.2, 7.7, 1.5),
              'Working title — the draft proposes one if blank.  → editor: Basics · Title', size=24),
            F('cover.training', 'Training it lives inside', (0.5, 2.9, 7.7, 1.0),
              'The course or program, with its code. Sets the register.  → wizard: The training it lives inside', size=14),
            F('cover.owner', 'Owner', (0.5, 4.1, 3.7, 1.0), 'Who to ask about this brief.', size=14),
            F('cover.version', 'Version · date', (4.5, 4.1, 3.7, 1.0), 'Your own versioning. PowerPoint keeps the history.', size=14),
            F('cover.status', 'Status', (0.5, 5.3, 7.7, 1.0),
              'Draft · SME review · Ready for the editor · Drafted in the editor', size=14),
        ]),
    'Brief': dict(
        title='THE BRIEF', chip='Wizard step 1 · The brief  →  Editor · Basics + Steps',
        fields=[
            F('brief.topic', 'Topic', (0.5, 1.2, 6.05, 1.25),
              'One line. Everything else builds on this.  → wizard: What is this scenario about?'),
            F('brief.learner', 'The learner plays', (6.78, 1.2, 6.05, 1.25),
              'Their position in the situation — placed to notice, decide or act.  → wizard: Who does the learner play?'),
            F('brief.steps', 'Steps, in order — one per line, each starting with its mode', (0.5, 2.7, 7.7, 4.35),
              'coach: / roleplay: / observe: then what happens, e.g. “coach: does this qualify as harassment?”  → wizard: The steps'),
            F('brief.sources', 'Source slides or documents', (8.43, 2.7, 4.4, 4.35),
              'Course file, slide numbers, policies. Paste the text on a SOURCE MATERIAL slide.'),
        ]),
    'Situation': dict(
        title='THE SITUATION', chip='Wizard step 2 · The world  →  Editor · Situation & world',
        fields=[
            F('situation.narrative', 'What’s true as the scenario opens — written to the learner', (0.5, 1.2, 7.55, 5.85),
              'Second person, present tense. The learner reads it, and it is the coach’s only picture of the setup.  → v4: narrative', size=11, verbatim=True),
            F('situation.setting', 'Setting', (8.28, 1.2, 4.55, 1.35),
              'Where this happens, in one line. The coach never sees it.  → v4: scene_world.setting'),
            F('situation.cast', 'Cast — Name — role — driver — never does', (8.28, 2.8, 4.55, 4.25),
              'Identity and disposition only; reactions go on the step.  → v4: scene_world.characters', size=11),
        ]),
    'Teaching': dict(
        title='THE TEACHING', chip='Wizard step 2 · The teaching  →  Editor · Teaching points + Coach voice',
        fields=[
            F('teaching.mustknows', 'Must-knows — 3 to 6, one per line', (0.5, 1.2, 6.05, 3.0),
              'What every learner walks away knowing — the teaching points and the expert answer.  → v4: teaching_points'),
            F('teaching.strongweak', 'Strong handling vs. weak', (0.5, 4.45, 6.05, 2.6),
              'What separates a strong pass from a thin one. Feeds the grading tiers on every step.  → wizard'),
            F('teaching.misconceptions', 'Common wrong answers → the redirect', (6.78, 1.2, 6.05, 3.85),
              'The misconception in the learner’s words, an arrow, then the coach’s redirect.  → v4: misconceptions'),
            F('teaching.voice', 'Coach voice', (6.78, 5.3, 6.05, 1.75),
              'How the coach comes across — a stance, not engine rules.  → v4: coach_persona + tone_guidelines'),
        ]),
    'Opening': dict(
        title='THE OPENING', chip='Optional warm-up · delete the slide to skip it  →  Editor · Opening reflection',
        fields=[
            F('opening.question', 'Warm-up question — verbatim, the coach’s first line', (0.5, 1.2, 8.4, 2.35),
              'One ungraded gut-reaction question before the first step. In quotes = the learner reads exactly this.  → v4: opening', size=14, verbatim=True),
            F('opening.turns', 'Turns', (9.13, 1.2, 3.7, 2.35),
              '1 or 2 learner turns. Calibrated, never graded.  → v4: opening.exit', size=14),
            F('opening.listen', 'Listen for — and how to acknowledge it', (0.5, 3.8, 12.33, 3.25),
              'Guidance for the coach: what a typical first answer contains (the misconception to expect), and how to acknowledge it without grading or previewing what comes next.  → v4: opening.levels'),
        ]),
    'Step': dict(
        title='STEP N · The story’s own name for this segment', chip='Duplicate this slide for each step  →  Editor · Steps',
        fields=[
            F('step.mode', 'Mode', (0.5, 1.2, 2.3, 0.72), 'coach · roleplay · observe', size=11),
            F('step.right', 'Right answer?', (2.96, 1.2, 2.3, 0.72), 'Yes = the coach lands it. No = open.', size=11),
            F('step.turns', 'Turns', (5.42, 1.2, 1.1, 0.72), 'Learner turns.', size=11),
            F('step.does', 'What the learner does here — and why now', (0.5, 2.02, 6.02, 1.05),
              'Guidance for the coach: the task, and what must be done to complete the step.  → v4: practice.purpose', size=10),
            F('step.opener', 'Opener — verbatim', (6.75, 1.2, 6.08, 1.87),
              'coach: the task line · roleplay: Narrator / Character lines · observe: the brief over the exhibit', size=10, verbatim=True),
        ] + tier_boxes(3.2, 2.75, ['Look for:', 'Respond:', 'Scene moves (roleplay only):'], 'step.tier',
                       'Recognise it · respond · move the scene  → v4: levels', size=9.5)
        + [
            F('step.debrief', 'Debrief lands — 2 to 4 points every learner hears, however the attempt went', (0.5, 6.05, 12.33, 1.03),
              'Guidance for the coach. One point per line, or separated by  ·   → v4: debrief.key_points', size=9.5),
        ]),
    'Examples': dict(
        title='STEP N · EXAMPLES (optional)', chip='Optional · one example per tier — the wizard drafts these if you don’t',
        fields=tier_boxes(1.2, 5.85, ['Learner:', 'Reply:'], 'example.tier',
                          'One learner line and the reply they hear.  → v4: levels.example', size=11)),
    'Close': dict(
        title='THE CLOSE', chip='Wizard · The expert answer  →  Editor · Expert answer',
        fields=[
            F('close.components', 'Ideal response components — a heading, then its points', (0.5, 1.2, 6.05, 5.85),
              'SME-validated statements; ships verbatim to every learner — the audit record.  → v4: component_groups', size=10.5),
            F('close.summary', 'Summary — verbatim', (6.78, 1.2, 6.05, 3.35),
              'Two to four sentences tying the components together — the last thing the learner reads.  → v4: summary', size=11, verbatim=True),
            F('close.authorities', 'External authorities', (6.78, 4.8, 2.9, 2.25),
              'Regulations, standards only. Ships to the learner.', size=10),
            F('close.internal', 'Internal source slides', (9.93, 4.8, 2.9, 2.25),
              'Course slide numbers. Never ships.', size=10),
        ]),
    'Source': dict(
        title='SOURCE MATERIAL', chip='Optional · paste anything — the wizard mines it for specifics',
        fields=[
            F('source.text', 'Source text', (0.5, 1.2, 12.33, 5.85),
              'The static scenario this replaces, the slide outline, a policy excerpt, SME notes. Add more slides if it runs long.  → wizard: Source material', size=10.5),
        ]),
}

HOW_IT_WORKS = [
    ('Type into the fields.', ' Every field is a placeholder. Outline view exports exactly what you type — and only that.'),
    ('Keep the label line', ' at the top of each field. It is how the editor knows what it is reading.'),
    ('One idea per line', ' in list fields. No bullets needed.'),
    ('Quotes mean verbatim.', ' A field marked verbatim is what the learner reads, word for word. Everything else is guidance for the coach.'),
    ('Duplicate a step slide', ' for each new step (⌘D). Don’t insert one from the layout gallery — it arrives without its labels.'),
    ('No engine notes.', ' Triggers, input types, partner labels and button behavior are fixed by the player.'),
    ('Copy out:', ' View → Outline → select all → paste into the editor’s Start-from-scratch source box.'),
]


# =============================================================================
# XML helpers
# =============================================================================
def clear_layout(layout):
    tree = layout.shapes._spTree
    for el in list(tree):
        if el.tag in (qn('p:sp'), qn('p:pic'), qn('p:grpSp'), qn('p:graphicFrame'), qn('p:cxnSp')):
            tree.remove(el)


def move_to_layout(shape, layout):
    layout.shapes._spTree.append(shape._element)


def renumber_ids(layout):
    n = 2
    for el in layout.shapes._spTree.iter(qn('p:cNvPr')):
        el.set('id', str(n))
        n += 1


def add_run(p, text, size, bold=False, italic=False, color=INK, spc=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if spc is not None:
        r._r.get_or_add_rPr().set('spc', str(spc))
    return r


def end_para_props(p, size, bold=False, color=INK):
    """What the author's typing inherits when they click on an empty line."""
    e = p._p.find(qn('a:endParaRPr'))
    if e is None:
        e = etree.SubElement(p._p, qn('a:endParaRPr'))
    e.set('sz', str(int(round(size * 100))))
    e.set('b', '1' if bold else '0')
    e.set('i', '0')
    for c in list(e):
        e.remove(c)
    sf = etree.SubElement(e, qn('a:solidFill'))
    etree.SubElement(sf, qn('a:srgbClr')).set('val', color)
    etree.SubElement(e, qn('a:latin')).set('typeface', FONT)


def panel(scratch, x, y, w, h, fill=TINT):
    shp = scratch.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.045
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ''
    shp.name = 'panel'
    return shp


def textbox(scratch, x, y, w, h, runs, size=9, italic=False, color=MUTED, anchor=MSO_ANCHOR.TOP, name='guide'):
    tb = scratch.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    for (txt, b) in runs:
        add_run(p, txt, size, bold=b, italic=italic and not b, color=color)
    return tb


def make_placeholder(scratch, layout, spec, idx, kind='body'):
    """A real placeholder on the LAYOUT, with a list style so slide text
    inherits our size/no-bullet defaults, and prompt text for empty use."""
    x, y, w, h = spec['box']
    if kind == 'body':
        h = h - HELP_H
    tb = scratch.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sp = tb._element
    sp.nvSpPr.cNvPr.set('name', spec['key'])
    sp.nvSpPr.cNvSpPr.attrib.pop('txBox', None)
    ph = etree.SubElement(sp.nvSpPr.nvPr, qn('p:ph'))
    if kind == 'title':
        ph.set('type', 'title')
    else:
        ph.set('type', 'body')
        ph.set('idx', str(idx))

    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.07)
    tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    size = spec['size']
    lst = sp.txBody.find(qn('a:lstStyle'))
    lvl1 = etree.SubElement(lst, qn('a:lvl1pPr'))
    lvl1.set('marL', '0')
    lvl1.set('indent', '0')
    spc = etree.SubElement(lvl1, qn('a:spcBef'))
    etree.SubElement(spc, qn('a:spcPts')).set('val', '200')
    etree.SubElement(lvl1, qn('a:buNone'))
    d = etree.SubElement(lvl1, qn('a:defRPr'))
    d.set('sz', str(int(round(size * 100))))
    d.set('b', '1' if kind == 'title' else '0')
    sf = etree.SubElement(d, qn('a:solidFill'))
    etree.SubElement(sf, qn('a:srgbClr')).set('val', INK)
    etree.SubElement(d, qn('a:latin')).set('typeface', FONT)
    if kind == 'body':
        lvl2 = etree.SubElement(lst, qn('a:lvl2pPr'))
        lvl2.set('marL', '228600')
        lvl2.set('indent', '-228600')
        bu = etree.SubElement(lvl2, qn('a:buChar'))
        bu.set('char', '•')
        d2 = etree.SubElement(lvl2, qn('a:defRPr'))
        d2.set('sz', str(int(round(size * 100))))
        etree.SubElement(d2, qn('a:latin')).set('typeface', FONT)

    # prompt text — shows only while the slide placeholder is empty
    p = tf.paragraphs[0]
    if kind == 'title':
        add_run(p, spec['label'], size, bold=True, color=INK)
    else:
        add_run(p, spec['label'].upper(), 9, bold=True, color=TIER.get(spec.get('tier'), BLUE), spc=60)
        p2 = tf.add_paragraph()
        add_run(p2, 'Type here — keep the label line above.', size, italic=True, color=MUTED)
    move_to_layout(tb, layout)
    return sp


def build_layout(layout, name, spec):
    """Turn one of the default template's layouts into ours."""
    clear_layout(layout)
    layout.name = name
    prs = layout.part.package.presentation_part.presentation
    scratch = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    for shp in list(scratch.shapes):
        shp._element.getparent().remove(shp._element)

    # decoration first (behind)
    for f in spec['fields']:
        x, y, w, h = f['box']
        move_to_layout(panel(scratch, x, y, w, h), layout)
    chip_w = 0.085 * len(spec['chip']) + 0.4
    chip_w = min(chip_w, 8.6)
    chip = panel(scratch, SLIDE_W - 0.5 - chip_w, 0.42, chip_w, 0.36, fill=TINT2)
    move_to_layout(chip, layout)
    move_to_layout(textbox(scratch, SLIDE_W - 0.5 - chip_w, 0.42, chip_w, 0.36, [(spec['chip'], False)],
                           size=9, color=INK2, anchor=MSO_ANCHOR.MIDDLE, name='chip'), layout)
    for f in spec['fields']:
        x, y, w, h = f['box']
        runs = [(f['help'], False)]
        if f.get('verbatim'):
            runs = [('VERBATIM  ', True)] + runs
        move_to_layout(textbox(scratch, x, y + h - HELP_H, w, HELP_H, runs, size=7.5, italic=True, color=MUTED,
                               anchor=MSO_ANCHOR.MIDDLE, name='help'), layout)
    move_to_layout(textbox(scratch, 0.5, 7.12, 6.5, 0.3, [('Scenario Brief · template v1', True), ('  ·  Aithera Scenario Simulator', False)],
                           size=8, color=MUTED, name='footer'), layout)
    move_to_layout(textbox(scratch, 6.9, 7.12, 5.93, 0.3, [('Fields are placeholders — type into them. Text boxes and tables are not exported.', False)],
                           size=8, italic=True, color=MUTED, name='footer'), layout)
    if name == 'Cover':
        # the how-to panel — layout-level, so it can't be edited or exported
        move_to_layout(panel(scratch, 8.6, 1.2, 4.23, 5.1, fill=TINT2), layout)
        tb = scratch.shapes.add_textbox(Inches(8.6), Inches(1.2), Inches(4.23), Inches(5.1))
        tb.name = 'howto'
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.16)
        p = tf.paragraphs[0]
        add_run(p, 'HOW THIS BRIEF WORKS', 9, bold=True, color=BLUE, spc=60)
        for head, rest in HOW_IT_WORKS:
            p = tf.add_paragraph()
            p.space_before = Pt(5)
            add_run(p, head, 10, bold=True, color=INK)
            add_run(p, rest, 10, color=INK2)
        move_to_layout(tb, layout)

    # placeholders on top
    make_placeholder(scratch, layout, dict(key='title', label=spec['title'], box=(0.5, 0.3, SLIDE_W - 1.0 - chip_w - 0.2, 0.7), size=24), 0, kind='title')
    for i, f in enumerate(spec['fields']):
        make_placeholder(scratch, layout, f, 10 + i)
    renumber_ids(layout)

    # drop the scratch slide
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[-1]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


# =============================================================================
# fit estimation — no renderer on this machine, so size text by arithmetic
# =============================================================================
CHAR_W = 0.50   # Calibri average advance as a fraction of the point size (conservative)
LINE_H = 1.22


def lines_needed(text, size, width_in):
    if not text:
        return 1
    cpl = max(8, int((width_in - 0.24) / (size / 72.0 * CHAR_W)))
    return sum(max(1, math.ceil(len(line) / cpl)) for line in text.split('\n'))


def fit_size(blocks, width_in, height_in, base, floor=8.0):
    """blocks: list of (text, size_or_None) — None means 'scale this one'."""
    size = base
    while True:
        h = 0.12
        for text, fixed in blocks:
            s = fixed if fixed else size
            h += lines_needed(text, s, width_in) * (s / 72.0 * LINE_H) + 0.028
        if h <= height_in or size <= floor:
            return size, h
        size -= 0.5


# =============================================================================
# filling slides
# =============================================================================
def slide_ph(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    raise KeyError(idx)


def write_value(tf, text, size):
    if not text:
        p = tf.add_paragraph()
        end_para_props(p, size)
        return
    for line in text.split('\n'):
        p = tf.add_paragraph()
        if line:
            add_run(p, line, size)
        end_para_props(p, size)


def fill_field(ph, spec, value, report):
    tf = ph.text_frame
    tf.clear()
    x, y, w, h = spec['box']
    inner_h = h - HELP_H
    label = spec['label'].upper()
    subs = spec.get('subs')
    blocks = [(label, 9)]
    if subs:
        for s in subs:
            blocks.append((s, 9.5))
            blocks.append(((value or {}).get(s, ''), None))
    else:
        blocks.append((value or '', None))
    size, est_h = fit_size(blocks, w, inner_h, spec['size'])
    report.append((spec['key'], size, est_h, inner_h))

    p0 = tf.paragraphs[0]
    add_run(p0, label, 9, bold=True, color=TIER.get(spec.get('tier'), BLUE), spc=60)
    end_para_props(p0, size)
    if subs:
        for s in subs:
            p = tf.add_paragraph()
            p.space_before = Pt(4)
            add_run(p, s, 9.5, bold=True, color=INK2)
            end_para_props(p, size)
            write_value(tf, (value or {}).get(s, ''), size)
    else:
        write_value(tf, value, size)


def add_from_layout(prs, layouts, name, title, values, report, title_size=None):
    slide = prs.slides.add_slide(layouts[name])
    spec = LAYOUTS[name]
    t = slide.shapes.title
    t.text_frame.clear()
    add_run(t.text_frame.paragraphs[0], title, title_size or 24, bold=True, color=INK)
    for i, f in enumerate(spec['fields']):
        ph = slide_ph(slide, 10 + i)
        ph._element.nvSpPr.cNvPr.set('name', f['key'])
        fill_field(ph, f, values.get(f['key']), report)
    return slide


def build_deck(path, content, n_steps, example_steps, include_source, notes):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    master = prs.slide_master
    defaults = list(master.slide_layouts)
    names = list(LAYOUTS.keys())
    layouts = {}
    for i, name in enumerate(names):
        build_layout(defaults[i], name, LAYOUTS[name])
        layouts[name] = defaults[i]
    for lay in defaults[len(names):]:
        master.slide_layouts.remove(lay)

    report = []
    cov = add_from_layout(prs, layouts, 'Cover', 'SCENARIO BRIEF', content.get('cover', {}), report)
    cov.notes_slide.notes_text_frame.text = notes
    add_from_layout(prs, layouts, 'Brief', 'THE BRIEF', content.get('brief', {}), report)
    add_from_layout(prs, layouts, 'Situation', 'THE SITUATION', content.get('situation', {}), report)
    add_from_layout(prs, layouts, 'Teaching', 'THE TEACHING', content.get('teaching', {}), report)
    add_from_layout(prs, layouts, 'Opening', 'THE OPENING', content.get('opening', {}), report)
    steps = content.get('steps', [])
    for n in range(1, n_steps + 1):
        st = steps[n - 1] if n - 1 < len(steps) else {}
        vals = {f'step.{k}': st.get(k) for k in ('mode', 'right', 'turns', 'does', 'opener', 'debrief')}
        for t in TIERS:
            vals[f'step.tier.{t}'] = (st.get('tiers') or {}).get(t)
        label = st.get('label') or 'The story’s own name for this segment'
        add_from_layout(prs, layouts, 'Step', f'STEP {n} · {label}', vals, report, title_size=20)
        if n in example_steps:
            ex = (st.get('examples') or {})
            vals = {f'example.tier.{t}': ex.get(t) for t in TIERS}
            add_from_layout(prs, layouts, 'Examples', f'STEP {n} · EXAMPLES (optional)', vals, report, title_size=20)
    add_from_layout(prs, layouts, 'Close', 'THE CLOSE', content.get('close', {}), report)
    if include_source:
        add_from_layout(prs, layouts, 'Source', 'SOURCE MATERIAL', content.get('source', {}), report)

    prs.core_properties.title = content.get('doc_title', 'Scenario Brief')
    prs.core_properties.author = 'Aithera UX · Scenario Simulator'
    prs.core_properties.subject = 'Scenario Brief — the K&A input to the Scenario Editor'
    prs.save(path)
    return report


# =============================================================================
# CONTENT — the Marshall worked example, re-homed from POC deck v3 (May 2026)
# =============================================================================
MARSHALL = dict(
    doc_title='Scenario Brief — The Marshall Scenario (worked example)',
    cover={
        'cover.title': 'The Marshall Scenario\nSex-Based Harassment, Empathy & Bystander Intervention',
        'cover.training': 'Harassment Prevention for Employees (JCOM-40198)',
        'cover.owner': 'K&A team · Scenario Simulator POC',
        'cover.version': 'Brief v1 · re-homed from POC deck v3 (May 2026) · 2026-09-11',
        'cover.status': 'Worked example — the 32-slide POC deck, restated on the Scenario Brief template',
    },
    brief={
        'brief.topic': 'Sex-based harassment, gender stereotyping and bystander intervention',
        'brief.learner': 'A co-worker who has observed the incidents involving Marshall — not his manager',
        'brief.steps': (
            'coach: Does this qualify as harassment? Reason it through under Title VII — there is a right answer, and it is yes.\n'
            'coach: What is Marshall experiencing? Set the law aside — the personal, professional and team-wide impact.\n'
            'roleplay: The break room — Jake makes a ‘Marsha’ comment with the room watching, and the learner has to act.'
        ),
        'brief.sources': (
            'JCOM-40198 Harassment Prevention for Employees Base Course.pptx\n'
            'Slides 18–24 (the Marshall learning object)\n'
            'Slide 13 (Title VII foundation)\n'
            'Slides 40–41 (Ira knowledge check — bystander guidance)'
        ),
    },
    situation={
        'situation.narrative': (
            'You’ve been working alongside Marshall for about eight months. He’s an administrative assistant — organized, a good communicator, clearly someone who takes his job seriously. But lately, he’s not himself.\n'
            '\n'
            'It started with Ethan, the project manager. He’d greet Marshall with ‘Hey Marsha!’ in the hallway. A couple of times he asked if Marshall had a skirt on ‘under that desk.’ Marshall let it go. He thought some joking might come with the job — especially given the way he dresses. So he tried not to make it a thing.\n'
            '\n'
            'Then Jake started. A junior engineer, hired not long after Marshall. He’d ask if the coffee was made whenever he passed Marshall’s desk. He’d refer to Marshall’s role as a ‘cozy lady job.’ What started as occasional became almost daily. The kind of remark that gets a few laughs and then everyone moves on — except Marshall doesn’t move on. He carries it.\n'
            '\n'
            'What Marshall didn’t know, not at first, was that there was a group chat. Someone eventually showed him: sexist memes, jokes. And two altered images — one with his face on a woman in a frilly princess dress, another with his face on a lingerie model’s body, captioned ‘Marsha’s true calling.’\n'
            '\n'
            'He was going to try to let it go. Until those images ended up on public social media — shareable, commentable, out there.\n'
            '\n'
            'Marshall has gotten quieter — he keeps his head down, doesn’t linger. You’ve just been made aware of the behaviors, and you’re not sure how to address them.'
        ),
        'situation.setting': 'An office workplace — an open floor of desks, a hallway, and a shared break room.',
        'situation.cast': (
            'Jake — junior engineer, hired not long after Marshall; the source of the daily ‘Marsha’ and ‘cozy lady job’ remarks — reads the room’s laughter as permission; a line held without heat makes him retreat to ‘it was a joke’ — never physical, never concedes it is harassment\n'
            '\n'
            'Marshall — administrative assistant, eight months in; the person the conduct is aimed at — responds to whether anyone in the room acknowledges what just happened — understated, a glance or a short line; never speaks for Jake, never asks the learner to step in for him\n'
            '\n'
            'Ethan — project manager, senior to both; started the ‘Marsha’ nickname — wants to be liked, hears a challenge as someone being humorless — never openly hostile; the harm arrives wrapped in friendliness'
        ),
    },
    teaching={
        'teaching.mustknows': (
            'Gender-stereotype-based conduct is sex-based harassment under Title VII — no explicit sexual advance or quid pro quo required.\n'
            'Hostile work environment: pervasive gender-based conduct that makes the workplace intimidating qualifies, and it affects everyone present, not only the target.\n'
            'Same-sex harassment is fully covered, and intent doesn’t matter — the test is impact and context.\n'
            'Sustained harassment causes documented psychological and career harm. ‘Just jokes’ is never an accurate frame.\n'
            'Marshall should report immediately — with documentation of incidents, dates and witnesses.\n'
            'Bystanders: Pick an Action · Offer Support · Consider Escalating — and check the organization’s policy for witness obligations.'
        ),
        'teaching.strongweak': (
            'Strong: names the gender-stereotype basis, applies the hostile-work-environment standard, arrives at reporting. In the break room, gives a clear direct or indirect signal, checks in with Marshall afterward, and considers escalating per policy.\n'
            '\n'
            'Weak: ‘not sexual, so not harassment’; ‘just banter’; ‘not my place.’ In the room: silence, laughing along, or putting the onus on Marshall to stand up for himself.'
        ),
        'teaching.misconceptions': (
            '‘It’s not sexual, so it’s not harassment.’ → The most common one. Title VII’s scope includes gender stereotyping — nothing sexual has to be said or demanded.\n'
            '‘It’s just joking / workplace banter.’ → Minimizes the cumulative weight. What does daily degradation actually cost someone?\n'
            '‘Not my place to get involved.’ → Silence is never neutral — it signals acceptance. Others are likely as concerned as you are and will support you acting.\n'
            '‘It only counts if it affects his job.’ → Misapplies the standard. A hostile work environment qualifies without economic injury.\n'
            '‘He knew some of this would come with the job, given how he dresses.’ → Anticipating mistreatment does not make it lawful; presentation is never consent.'
        ),
        'teaching.voice': (
            'Precise but not clinical — names the legal reality without lecturing. Affirms before it corrects: meets the learner’s instinct, then sharpens it. Never shames a bystander instinct — redirects with curiosity and specificity. Ends every response with a question or a forward pivot.'
        ),
    },
    opening={
        'opening.question': '“Before we get into the specifics — take a moment. What’s your gut reaction to this behavior? Is anything about this situation standing out to you, or feeling unclear?”',
        'opening.turns': '1',
        'opening.listen': (
            'Starting assumptions — ‘nothing sexual is really happening’, ‘it seems mean but not harassment’ — and how comfortable they are with getting involved at all.\n'
            'Acknowledge what they said and reflect a little of it back. Do not grade it and do not preview the law. Then move on: “Now let’s take a closer look at what’s actually happening here.”'
        ),
    },
    steps=[
        dict(
            label='Does This Qualify as Harassment?',
            mode='coach',
            right='Yes — the coach lands it plainly.',
            turns='2',
            does='Reason through whether the conduct qualifies as sexual harassment under Title VII. One probing follow-up at most, no teaching; then the coach lands the legal conclusion plainly, whatever the tier.',
            opener='“Based on what you know about workplace harassment — think through what Marshall is experiencing. In your view, does this qualify as sexual harassment? Walk through your reasoning.”',
            tiers={
                'unthoughtful': {
                    'Look for:': 'Conflates harassment with explicit sexual advances or quid pro quo. May suggest Marshall’s dress, or his expectation that ‘some joking would come with the job’, reduces the severity. Calls it teasing or bullying. Misses the gender-stereotype basis entirely.',
                    'Respond:': 'Clarify the two types of harassment. Address the ‘he knew it would happen / the way he dresses’ framing directly: anticipating mistreatment does not make it legal, and presentation is not consent. Deliver the answer — yes, this is harassment, and Marshall should report it.',
                    'Scene moves (roleplay only):': '',
                },
                'neutral': {
                    'Look for:': 'Recognizes the conduct feels targeted and wrong; intuits the gender-stereotyping angle. Stuck on the quid pro quo model — believes an exchange or direct threat is required. Unsure whether it ‘technically’ qualifies.',
                    'Respond:': 'Affirm the gender-targeting observation. Distinguish quid pro quo from hostile work environment: pervasive gender-based conduct that makes the workplace intimidating qualifies — no exchange required. Confirm: this is harassment under Title VII and Marshall should report it now.',
                    'Scene moves (roleplay only):': '',
                },
                'strong': {
                    'Look for:': 'Correctly identifies gender stereotyping as the basis. Applies the hostile-work-environment standard. Notes harassment need not be explicitly sexual. Arrives at reporting. May note same-sex protections.',
                    'Respond:': 'Validate the full analysis and confirm the grounding: Title VII, hostile work environment, gender stereotyping. Add the same-sex point if not raised. Reinforce that the public images are a significant escalation — reporting should be urgent and documented.',
                    'Scene moves (roleplay only):': '',
                },
            },
            debrief=(
                'There are two types of harassment — quid pro quo and hostile work environment — and no exchange or threat is required for the second.  ·  '
                'Gender-stereotype conduct is sex-based harassment under Title VII; how someone presents is never consent.  ·  '
                'The public images are a major escalation — report, documented, now.'
            ),
        ),
        dict(
            label='What Is Marshall Experiencing?',
            mode='coach',
            right='No — open judgment.',
            turns='2',
            does='Set the law aside and think about Marshall as a person: what this is doing to him professionally and personally, and to everyone around him. One probing follow-up at most before the coach pulls it together.',
            opener='“Now that we’ve established what this is legally, let’s shift perspective. Set the legal framework aside for a moment and think about Marshall as a person. What do you think this situation is doing to him professionally and personally? How could this affect others in your workplace?”',
            tiers={
                'unthoughtful': {
                    'Look for:': 'Minimizes the impact as surface-level embarrassment. ‘Just jokes’, ‘brush it off.’ Treats Marshall’s reaction as a matter of personal resilience. Doesn’t see the cumulative weight of daily, sustained harassment.',
                    'Respond:': 'Gently challenge the ‘brush it off’ frame. Validate that resilience is real, then introduce what research shows: sustained harassment links to anxiety, performance decline and loss of motivation. Ask what it actually costs Marshall to keep ‘staying professional’ every day.',
                    'Scene moves (roleplay only):': '',
                },
                'neutral': {
                    'Look for:': 'Recognizes psychological discomfort — anxiety, dread, reluctance to be visible. Correctly reads the public images as an escalation. Stays surface-level on the career dimension.',
                    'Respond:': 'Affirm the psychological read and the public-image point. Extend to the career dimension: eight months in is a critical window for credibility. What does it mean to spend energy managing humiliation instead of doing your best work?',
                    'Scene moves (roleplay only):': '',
                },
                'strong': {
                    'Look for:': 'Genuine empathy. Names cumulative, dignity-level harm. Connects personal and professional impact. Reads the public images as serious escalation. May name the power dynamic — a project manager is a participant.',
                    'Respond:': 'Validate fully. Add the broader dimension: unchallenged conduct resets what feels normal and who feels safe to speak up — that is what a hostile work environment means in practice. Every time someone sees this and says nothing, Marshall learns he is alone. That is the bridge to the break room.',
                    'Scene moves (roleplay only):': '',
                },
            },
            debrief=(
                'Sustained harassment links to anxiety, declining performance and lost motivation — ‘putting it aside’ has a real cost.  ·  '
                'Eight months in is a critical window for credibility; energy spent managing humiliation is energy not spent on the work.  ·  '
                'Unchallenged conduct resets what the whole team treats as normal — which is exactly where bystanders come in.'
            ),
        ),
        dict(
            label='Bystander Intervention — The Break Room',
            mode='roleplay — Jake',
            right='No — open judgment.',
            turns='2',
            does='Walk in as Jake makes a ‘Marsha’ comment with the room watching, and act — what the learner types is what they do and say in front of everyone. The scene reacts to the approach; after turn two it closes and the coach debriefs.',
            opener=(
                'Narrator: Marshall is getting coffee and a few others are sitting around talking. Jake stands next to Marshall and pours himself a cup.\n'
                'Jake: “Hey, did you make this? Guess that’s what you’re here for — living your best Marsha life.”\n'
                'Narrator: He grins and looks around as you walk into the break room and witness the exchange.'
            ),
            tiers={
                'unthoughtful': {
                    'Look for:': 'Non-intervention — looks away, stays quiet, laughs along; ‘not my place.’ Or aggression toward Jake. Or putting the onus on Marshall (‘stand up for yourself’).',
                    'Respond:': 'Name what silence communicates: Jake reads permission, Marshall reads that no one sees it. Others are likely uncomfortable too and will back you. Intervening doesn’t require confrontation — offer low-key options, ask what feels manageable.',
                    'Scene moves (roleplay only):': 'Jake reads silence as a green light and escalates for laughs, or meets aggression in kind. Marshall stares at the floor, put on the spot or unsupported.',
                },
                'neutral': {
                    'Look for:': 'Uncomfortable and wants to act — a look, a subject change, a vague redirect. Right instinct, vague execution: nothing clearly signals the behavior is a problem. Unlikely to check in with Marshall after.',
                    'Respond:': 'Affirm the redirect as a real bystander move, then push for specificity — ‘shift things somehow’ stays invisible. Offer a concrete pivot (‘Hey Jake, what’s the update on Henderson?’) and add the follow-up with Marshall.',
                    'Scene moves (roleplay only):': 'Jake breezes past the redirect or excuses it as a joke; the subject change doesn’t take. Marshall glances at the learner — unclear what he made of it.',
                },
                'strong': {
                    'Look for:': 'A direct or indirect signal that the behavior isn’t supported — ‘I don’t really appreciate that’, ‘that’s not a professional thing to say.’ Plans to check in with Marshall privately. Considers escalating per policy.',
                    'Respond:': 'Validate all three components — action in the moment, follow-up with Marshall, escalation. Note that the organization’s harassment policy may define specific witness obligations. Any of the three moves beats silence.',
                    'Scene moves (roleplay only):': 'The callout lands, but Jake resists — ‘Whoa, relax, it was a joke. Right, Marshall?’ — and glances around for backup. The room goes quiet. Marshall is steadier, not alone.',
                },
            },
            debrief=(
                'Silence is not neutral — Jake reads it as permission and Marshall reads it as no one noticing.  ·  '
                'Pick an Action: a direct signal or an indirect redirect both count; confrontation is one option, not the only one.  ·  '
                'Offer Support: check in with Marshall privately after.  ·  '
                'Consider Escalating: the organization’s policy may define witness obligations, and a bystander can report independently of Marshall.'
            ),
            examples={
                'unthoughtful': {
                    'Learner:': '“Honestly I’d probably just look down. I don’t want to get in the middle of something between two colleagues. If I say something, I might make it a bigger deal than it needs to be.”',
                    'Reply:': 'Narrator: You keep your eyes on the floor and stay quiet. Jake reads the silence as a green light.\nJake: “Anyone else need dear little Marsha to make them a cup of coffee?”\nNarrator: A couple of people chuckle along and Marshall stares at the floor, edging toward the door. The moment is still hanging there.',
                },
                'neutral': {
                    'Learner:': '“I might give Jake a look or try to change the subject quickly. I wouldn’t call it out directly but I’d try to shift things somehow.”',
                    'Reply:': 'Narrator: You catch Jake’s eye and try to steer things elsewhere, but Jake breezes right past it.\nJake: “I’m just saying, Marsha’s the best coffee maker we have.”\nNarrator: The subject-change doesn’t take; Jake loops back to the joke and the room is still watching. Marshall shifts, unsure whether anyone’s actually going to say something.',
                },
                'strong': {
                    'Learner:': '“I’d say ‘Hey, that’s not cool, Jake’ — not a confrontation, just enough to signal it’s not okay. Then I’d check in with Marshall after. And given how often this seems to happen, I’d probably talk to my manager or HR about it.”',
                    'Reply:': 'Narrator: You spoke up — and Jake doesn’t just let it go. His grin tightens.\nJake: “Whoa, relax, it was a joke. Right, Marshall? Tell them you’re not offended.”\nNarrator: He puts Marshall on the spot and glances around for backup. The room goes quiet, watching to see what you’ll do.',
                },
            },
        ),
    ],
    close={
        'close.components': (
            'What actually qualifies\n'
            'Gender stereotyping is sex-based harassment under Title VII — explicit sexual advances are not required.\n'
            'Hostile work environment standard — pervasive, gender-based conduct qualifies, and affects everyone in the environment, not only the primary target.\n'
            'Same-sex harassment is fully covered — the gender of harasser and target is irrelevant.\n'
            'Intent doesn’t determine harassment — the test is impact and context.\n'
            '\n'
            'Cumulative weight and reporting\n'
            'Cumulative weight is real — sustained harassment causes documented psychological and career harm, and affects others in the workplace.\n'
            'Marshall should report immediately — with documentation of incidents, dates and witnesses.\n'
            '\n'
            'The bystander framework\n'
            'Pick an action — a direct or indirect in-the-moment signal; direct confrontation is one option, not the only one. Others will support intervention.\n'
            'Offer support to the targeted person privately — follow up after the moment passes.\n'
            'Consider escalating — check your organization’s harassment policy for witness obligations; bystanders can report independently of the target.'
        ),
        'close.summary': (
            '“What happened to Marshall is sexual harassment. Gender stereotyping — mocking someone for not conforming to expectations about how a man should act — is a form of sex-based harassment under Title VII. It does not require sexual advances, and the intent behind the behavior is less important than its impact. Repeated conduct like this can create a hostile work environment, so it is important to document the incidents and report them. If you witness harassment, speak up if it’s safe, support the person affected, and follow your organization’s reporting procedures. Even a small action can make a meaningful difference.”'
        ),
        'close.authorities': 'Title VII of the Civil Rights Act — sex-based harassment, including conduct based on gender stereotypes, as sex discrimination.',
        'close.internal': (
            'Slide 19 — scenario narration\n'
            'Slide 20 — knowledge check: should Marshall report?\n'
            'Slide 21 — Marshall’s real case\n'
            'Slide 24 — hostile work environment definition\n'
            'Slide 13 — Title VII foundation\n'
            'Slides 40–41 — bystander guidance'
        ),
    },
)

NOTES_TEMPLATE = (
    'SCENARIO BRIEF — how to use this deck\n\n'
    'Every field is a PowerPoint placeholder. Type into the fields and keep the small label line at the top of each one. '
    'Do not add text boxes or tables — Outline view does not export them, so the editor would never see them.\n\n'
    'One slide per step: duplicate a STEP slide (Cmd/Ctrl+D) for each new step rather than inserting from the layout gallery. '
    'Delete THE OPENING if the scenario has no warm-up. Delete EXAMPLES slides you don’t need — the wizard drafts examples itself.\n\n'
    'Verbatim fields are what the learner reads, word for word. Every other field is guidance for the coach.\n\n'
    'To hand off: View → Outline → select all → copy → paste into the Scenario Editor’s Start-from-scratch source box. '
    'The wizard reads the labels and pre-fills its interview; the AI drafts only what you left blank.'
)


def main():
    os.makedirs(OUT, exist_ok=True)
    blank = dict(doc_title='Scenario Brief — template')
    r1 = build_deck(os.path.join(OUT, 'scenario-brief-template.pptx'), blank, n_steps=3, example_steps={1}, include_source=True, notes=NOTES_TEMPLATE)
    r2 = build_deck(os.path.join(OUT, 'scenario-brief-example-marshall.pptx'), MARSHALL, n_steps=3, example_steps={3}, include_source=False, notes=NOTES_TEMPLATE)
    print('wrote', OUT)
    tight = [(k, s, round(h, 2), round(a, 2)) for (k, s, h, a) in r2 if h > a * 0.97]
    print('example deck — fields sized at the floor or near capacity:')
    for row in tight:
        print('  ', row)
    return 0


if __name__ == '__main__':
    sys.exit(main())
