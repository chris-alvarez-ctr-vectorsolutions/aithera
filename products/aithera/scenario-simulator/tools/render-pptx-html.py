#!/usr/bin/env python3
"""
Render a python-pptx-built deck to a single HTML page of absolutely positioned
slides — a stand-in for PowerPoint on a machine with no Office/LibreOffice.

    python3 tools/render-pptx-html.py deck.pptx out.html

It draws rounded rectangles and text boxes (slide shapes over layout shapes),
honouring position, size, fill, font size/weight/italic/colour and insets.
Empty slide placeholders show nothing (as in a slideshow). Text wraps with the
browser's metrics, so treat overflow as approximate — the generator's own fit
arithmetic is the authority; this is for eyeballing layout, density and gaps.
"""
import html
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu

EMU_IN = 914400.0


def css_color(font):
    try:
        if font.color and font.color.type is not None and font.color.rgb is not None:
            return '#' + str(font.color.rgb)
    except Exception:
        pass
    return None


def fill_color(shape):
    try:
        if shape.fill.type == 1:  # solid
            return '#' + str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def ph_fill(shape, layout_shape):
    return fill_color(shape) or (fill_color(layout_shape) if layout_shape is not None else None)


def run_html(run, default_sz):
    f = run.font
    sz = f.size.pt if f.size else default_sz
    st = [f'font-size:{sz}pt']
    if f.bold:
        st.append('font-weight:700')
    if f.italic:
        st.append('font-style:italic')
    c = css_color(f)
    if c:
        st.append(f'color:{c}')
    rpr = run._r.find(qn('a:rPr'))
    if rpr is not None and rpr.get('spc'):
        st.append(f"letter-spacing:{int(rpr.get('spc'))/100.0}pt")
    return f'<span style="{";".join(st)}">{html.escape(run.text)}</span>'


def default_size(shape, layout_ph):
    """lvl1 defRPr size from the slide shape's or its layout placeholder's lstStyle."""
    for sp in (shape, layout_ph):
        if sp is None:
            continue
        try:
            d = sp._element.txBody.find(qn('a:lstStyle'))
            if d is None:
                continue
            l1 = d.find(qn('a:lvl1pPr'))
            if l1 is None:
                continue
            dr = l1.find(qn('a:defRPr'))
            if dr is not None and dr.get('sz'):
                return int(dr.get('sz')) / 100.0
        except Exception:
            pass
    return 18.0


def shape_html(shape, layout_ph=None, skip_empty_ph=False):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return ''
    geo = layout_ph if (layout_ph is not None and shape.left is None) else shape
    if geo.left is None:
        return ''
    x, y, w, h = (Emu(v).inches for v in (geo.left, geo.top, geo.width, geo.height))
    bg = ph_fill(shape, layout_ph)
    radius = 0
    try:
        if shape.auto_shape_type is not None and 'ROUNDED' in str(shape.auto_shape_type):
            radius = 0.045 * min(w, h) * 96 * 1.6
    except Exception:
        pass
    style = [f'left:{x}in', f'top:{y}in', f'width:{w}in', f'height:{h}in']
    if bg:
        style.append(f'background:{bg}')
    if radius:
        style.append(f'border-radius:{radius}px')
    inner = ''
    if shape.has_text_frame:
        tf = shape.text_frame
        paras = tf.paragraphs
        if skip_empty_ph and not ''.join(r.text for p in paras for r in p.runs).strip():
            return ''
        src = tf if shape.left is not None else (layout_ph.text_frame if layout_ph is not None else tf)
        ml = Emu(src.margin_left).inches if src.margin_left is not None else 0.1
        mr = Emu(src.margin_right).inches if src.margin_right is not None else 0.1
        mt = Emu(src.margin_top).inches if src.margin_top is not None else 0.05
        anchor = str(src.vertical_anchor or '')
        dsz = default_size(shape, layout_ph)
        ps = []
        for p in paras:
            sb = p.space_before.pt if p.space_before else 0
            body = ''.join(run_html(r, dsz) for r in p.runs) or '&nbsp;'
            # A paragraph's own font-size drives its line-height — a run's explicit
            # size only styles its <span>, so a short line (e.g. one run smaller
            # than the shape's default) would otherwise inherit the container's
            # larger default line-height and get clipped by the shape's overflow.
            para_sz = max((r.font.size.pt for r in p.runs if r.font.size), default=dsz)
            ps.append(f'<p style="margin:{sb}pt 0 0;font-size:{para_sz}pt">{body}</p>')
        va = 'center' if 'MIDDLE' in anchor else 'flex-start'
        inner = (f'<div class="tf" style="padding:{mt}in {mr}in 0 {ml}in;justify-content:{va};font-size:{dsz}pt">'
                 + ''.join(ps) + '</div>')
    return f'<div class="shp" style="{";".join(style)}">{inner}</div>'


def main(src, out):
    prs = Presentation(src)
    W = Emu(prs.slide_width).inches
    H = Emu(prs.slide_height).inches
    slides_html = []
    for n, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout
        lph = {}
        for sp in layout.placeholders:
            lph[sp.placeholder_format.idx] = sp
        parts = []
        for sp in layout.shapes:
            if sp.is_placeholder:
                continue
            parts.append(shape_html(sp))
        for sp in slide.shapes:
            lp = lph.get(sp.placeholder_format.idx) if sp.is_placeholder else None
            parts.append(shape_html(sp, lp, skip_empty_ph=True))
        slides_html.append(f'<section class="slide" style="width:{W}in;height:{H}in"><span class="n">{n} · {html.escape(layout.name)}</span>{"".join(parts)}</section>')
    doc = f'''<!doctype html><meta charset="utf-8"><title>render · {html.escape(src.split('/')[-1])}</title>
<style>
body{{margin:0;background:#3a3f47;font-family:Calibri,Carlito,"Helvetica Neue",Arial,sans-serif;line-height:1.2}}
.slide{{position:relative;background:#fff;margin:24px auto;box-shadow:0 6px 30px rgba(0,0,0,.4);overflow:hidden}}
.n{{position:absolute;right:6px;top:-20px;color:#ddd;font:11px monospace}}
.shp{{position:absolute;box-sizing:border-box;overflow:hidden}}
.tf{{display:flex;flex-direction:column;height:100%;box-sizing:border-box;word-wrap:break-word}}
.tf p{{white-space:pre-wrap}}
</style>{"".join(slides_html)}
<script>
// ?slide=N shows only that slide, flush to the viewport — for headless screenshots.
(function(){{var m=/[?&]slide=(\d+)/.exec(location.search); if(!m) return; var n=+m[1];
document.body.style.background='#fff';
document.querySelectorAll('.slide').forEach(function(s,i){{ if(i+1!==n){{s.remove();}} else {{s.style.margin='0';s.style.boxShadow='none';}} }});
var t=document.querySelector('.n'); if(t) t.remove();}})();
</script>'''
    open(out, 'w').write(doc)
    print('wrote', out, len(slides_html), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
